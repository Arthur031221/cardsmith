"""Shared test fixtures. All sample documents are built at runtime, nothing is stored."""

from __future__ import annotations

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches

CELL_TEXT = (
    "The mitochondrion is the organelle that converts nutrients into ATP through "
    "cellular respiration. Plant cells contain chloroplasts, which capture light "
    "energy during photosynthesis. The nucleus stores the cell's genetic material "
    "and controls protein synthesis by transcribing DNA into messenger RNA."
)


class FakeChatClient:
    """Stand-in for OllamaClient. Never touches the network."""

    def __init__(self, responses=None, script=None, model="fake-model"):
        self.calls: list[tuple[list[dict], dict]] = []
        self._responses = list(responses) if responses is not None else None
        self._script = script
        self.model = model

    def chat_json(self, messages, schema):
        self.calls.append((messages, schema))
        if self._script is not None:
            return self._script(messages, schema)
        if self._responses:
            return self._responses.pop(0)
        raise AssertionError("FakeChatClient: no scripted response left")

    def ping(self):
        return True, "ok"


def card_response(quote: str, n: int = 1) -> dict:
    """Build a valid CARD_RESPONSE_SCHEMA payload with a grounded quote."""
    cards = []
    for i in range(n):
        cards.append(
            {
                "type": "basic",
                "front": f"Question {i}?",
                "back": f"Answer {i}",
                "cloze_text": "",
                "source_quote": quote,
            }
        )
    return {"cards": cards}


@pytest.fixture
def fake_client():
    return FakeChatClient()


@pytest.fixture
def tmp_db(tmp_path) -> str:
    return str(tmp_path / "cardsmith-test.db")


@pytest.fixture
def sample_pdf(tmp_path) -> str:
    path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page()
        text = f"Page {i + 1}.\n{CELL_TEXT}"
        page.insert_text((72, 72), text, fontsize=11)
    doc.save(str(path))
    doc.close()
    return str(path)


@pytest.fixture
def sample_pptx(tmp_path) -> str:
    path = tmp_path / "sample.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for i in range(2):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Slide {i + 1}"
        body = slide.placeholders[1]
        body.text_frame.text = CELL_TEXT
    prs.slide_width = Inches(10)
    prs.save(str(path))
    return str(path)


@pytest.fixture
def sample_txt(tmp_path) -> str:
    path = tmp_path / "sample.txt"
    path.write_text(f"{CELL_TEXT}\n\n{CELL_TEXT[::-1][:50]} more notes here to pad it out.\n")
    return str(path)


@pytest.fixture
def empty_pdf(tmp_path) -> str:
    path = tmp_path / "empty.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return str(path)
