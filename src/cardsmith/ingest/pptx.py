"""Slide deck text extraction through python-pptx."""

from __future__ import annotations

from pptx import Presentation

from cardsmith.ingest.chunking import Chunk, merge_units


def extract_pptx(path: str) -> list[Chunk]:
    """Extract text per slide (title, body shapes, and speaker notes) and merge into chunks."""
    prs = Presentation(path)
    units: list[tuple[str, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        units.append((f"slide {i}", "\n".join(parts)))
    return merge_units(units)
